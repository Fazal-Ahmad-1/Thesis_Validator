"""
Step 4A-1 reliability test: correctly-formatted documents should pass.

For each of the three rule profiles, this builds an in-memory DOCX that
follows that profile's actual rules as closely as possible, then asserts:
    - is_valid is True
    - blocking_error_count is 0
    - score is 100

Run from backend/ with the venv active:
    python test_reliability_manual.py

Manual test only — no pytest, no new dependencies. Does not modify any
validator or rule file.

KNOWN CAVEAT (discovered while writing this test, not fixed here):
font_validator.py checks every paragraph in the document against the
profile's single body-font rule, including heading-styled paragraphs.
It does not exclude headings. The Amity fixture below intentionally
does NOT set an explicit font family on heading runs (only size/bold,
which is what heading_validator actually checks), so font_validator's
family resolution correctly falls through to "unresolved -> skip"
rather than flagging it. If a heading style DID explicitly declare a
font family matching the body font (a very plausible real-world case),
font_validator would report a false "wrong font size" violation for
that heading, since heading and body checks are not reconciled. This
is a real gap worth tracking, not a mistake in this fixture.
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from validator.rule_loader import load_rules
from validator.validator import validate_document


# --- OOXML helpers, reused from create_valid_asu_sample.py ---

def add_page_field(paragraph):
    """Insert a real Word PAGE field (simple field) into a paragraph."""
    p = paragraph._p
    fld_simple = OxmlElement("w:fldSimple")
    fld_simple.set(qn("w:instr"), "PAGE")
    run_el = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = "1"
    run_el.append(text_el)
    fld_simple.append(run_el)
    p.append(fld_simple)


def set_page_num_type(section, fmt: str, start: int):
    """Explicitly declare a section's page-numbering format/start (w:pgNumType)."""
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:fmt"), fmt)
    pg_num_type.set(qn("w:start"), str(start))


def add_body_paragraph(document, text, family, size_pt, line_spacing):
    """Add a body paragraph with explicit font and line spacing (never inherited)."""
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = family
    run.font.size = Pt(size_pt)
    paragraph.paragraph_format.line_spacing = line_spacing
    return paragraph


def add_heading(document, style_name, text, size_pt, bold, uppercase_text):
    """Add a heading paragraph with explicit size/bold (no font family set)."""
    paragraph = document.add_paragraph(style=style_name)
    run = paragraph.add_run(text.upper() if uppercase_text else text)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    return paragraph


# --- Per-profile fixture builders ---

def build_asu_document():
    document = Document()
    section = document.sections[0]

    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)

    # ASU: Times New Roman 12pt is in the "allowed" list; double spacing.
    add_body_paragraph(document, "Chapter 1: Introduction to the Study.", "Times New Roman", 12, 2.0)
    add_body_paragraph(document, "This paragraph discusses the background of the research.", "Times New Roman", 12, 2.0)

    # Pagination: PAGE field, bottom-center footer; explicit decimal/start=1.
    set_page_num_type(section, fmt="decimal", start=1)
    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer_paragraph)

    return document


def build_mumbai_document():
    document = Document()
    section = document.sections[0]

    # A4 size; margins from profile (mm): top 15, bottom 22, left 30, right 20.
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(30)
    section.right_margin = Mm(20)

    # Mumbai: Times New Roman 12pt; 1.5 line spacing.
    add_body_paragraph(document, "Chapter 1: Introduction.", "Times New Roman", 12, 1.5)
    add_body_paragraph(document, "This paragraph describes the methodology used.", "Times New Roman", 12, 1.5)

    # Pagination: PAGE field required, but no position/alignment specified
    # for this profile, so the field can be in either header or footer.
    set_page_num_type(section, fmt="decimal", start=1)
    footer_paragraph = section.footer.paragraphs[0]
    add_page_field(footer_paragraph)

    return document


def build_amity_document():
    document = Document()
    section = document.sections[0]

    # A4 size; margins 2.5cm on all sides.
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # Amity: Times New Roman 12pt body text; 1.5 line spacing.
    add_body_paragraph(document, "This paragraph is the body of the report.", "Times New Roman", 12, 1.5)

    # Amity headings: Heading 1 (14pt, bold, uppercase), Heading 2
    # (12pt, bold, uppercase), Heading 3 (12pt, bold). See module
    # docstring for why font family is deliberately left unset here.
    add_heading(document, "Heading 1", "Main Heading", size_pt=14, bold=True, uppercase_text=True)
    add_heading(document, "Heading 2", "Category Heading", size_pt=12, bold=True, uppercase_text=True)
    add_heading(document, "Heading 3", "Sub Heading", size_pt=12, bold=True, uppercase_text=False)

    # Pagination: PAGE field required, bottom-center footer. No
    # body_format/body_starts_at/preliminary_format for this profile,
    # so no pgNumType declaration is needed.
    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer_paragraph)

    return document


BUILDERS = {
    "asu_graduate": build_asu_document,
    "amity_selected": build_amity_document,
    "mumbai_engineering": build_mumbai_document,
}


def run_case(profile_name):
    builder = BUILDERS[profile_name]
    document = builder()

    tmp_path = f"_tmp_{profile_name}.docx"
    document.save(tmp_path)

    result = validate_document(tmp_path, profile_name)

    expected_is_valid = True
    expected_blocking_count = 0
    expected_score = 100

    passed = (
        result["is_valid"] == expected_is_valid
        and result["blocking_error_count"] == expected_blocking_count
        and result["score"] == expected_score
    )

    status = "PASS" if passed else "FAIL"
    print(f"[{status}] {profile_name}")
    print(f"    is_valid={result['is_valid']} (expected {expected_is_valid})")
    print(f"    blocking_error_count={result['blocking_error_count']} (expected {expected_blocking_count})")
    print(f"    score={result['score']} (expected {expected_score})")
    print(f"    informational_notice_count={result['informational_notice_count']}")

    if result["blocking_errors"]:
        print("    Blocking errors found:")
        for e in result["blocking_errors"]:
            print(f"      - [{e['category']}] {e['rule']}: {e['message']}")

    print()
    return passed


if __name__ == "__main__":
    print("Step 4A-1: correctly-formatted document reliability check\n")

    results = {profile: run_case(profile) for profile in BUILDERS}

    print("=" * 60)
    for profile, passed in results.items():
        print(f"{profile}: {'PASS' if passed else 'FAIL'}")