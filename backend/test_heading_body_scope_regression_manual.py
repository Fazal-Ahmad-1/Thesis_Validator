from docx import Document
from docx.shared import Pt

from validator.font_validator import validate_font_settings
from validator.paragraph_validator import validate_paragraph_settings
from validator.heading_validator import validate_heading_settings
from validator.rule_loader import load_rules


def make_doc(style_name="Heading 1", heading_size=14, heading_bold=True, heading_spacing=1.15, heading_text="A MAIN HEADING"):
    doc = Document()
    styles = doc.styles
    style = styles[style_name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(heading_size)
    style.font.bold = heading_bold
    style.paragraph_format.line_spacing = heading_spacing
    p = doc.add_paragraph(heading_text, style=style_name)
    body = doc.add_paragraph("Normal body text")
    body.runs[0].font.name = "Times New Roman"
    body.runs[0].font.size = Pt(12)
    body.paragraph_format.line_spacing = 1.5
    return doc


def check(label, condition):
    print(("[PASS] " if condition else "[FAIL] ") + label)
    return condition

all_ok = True

# UOK: configured Heading 1 must be excluded from body font/spacing checks.
uok = load_rules("uok_phd")
doc = make_doc(heading_size=14, heading_spacing=1.15)
font_errors = validate_font_settings(doc, uok)
spacing_errors = validate_paragraph_settings(doc, uok)
heading_errors = validate_heading_settings(doc, uok)
all_ok &= check("UOK configured heading does not create body-font error", len(font_errors) == 0)
all_ok &= check("UOK configured heading does not create body-spacing error", len(spacing_errors) == 0)
all_ok &= check("UOK correctly formatted heading has no heading error", len(heading_errors) == 0)

# UOK: wrong heading size should be caught by heading validator, not body font validator.
doc = make_doc(heading_size=12, heading_spacing=1.5)
font_errors = validate_font_settings(doc, uok)
spacing_errors = validate_paragraph_settings(doc, uok)
heading_errors = validate_heading_settings(doc, uok)
all_ok &= check("UOK wrong heading size is not reported as body-font error", len(font_errors) == 0)
all_ok &= check("UOK wrong heading spacing is not reported as body-spacing error", len(spacing_errors) == 0)
all_ok &= check("UOK wrong heading size is caught by heading validator", any(e.get("rule") == "Main Heading Size" for e in heading_errors))

# GGU: configured Heading 1 must also be excluded from body checks.
ggu = load_rules("ggu_political_science")
doc = make_doc(heading_size=12, heading_spacing=1.15)
font_errors = validate_font_settings(doc, ggu)
spacing_errors = validate_paragraph_settings(doc, ggu)
heading_errors = validate_heading_settings(doc, ggu)
all_ok &= check("GGU configured heading does not create body-font error", len(font_errors) == 0)
all_ok &= check("GGU configured heading does not create body-spacing error", len(spacing_errors) == 0)
all_ok &= check("GGU correctly formatted heading has no heading error", len(heading_errors) == 0)

# ASU has no configured heading rules: Heading 1 must NOT be automatically excluded.
asu = load_rules("asu_graduate")
doc = make_doc(heading_size=14, heading_spacing=1.15)
font_errors = validate_font_settings(doc, asu)
spacing_errors = validate_paragraph_settings(doc, asu)
all_ok &= check("ASU does not auto-exclude unconfigured Heading 1 from body-font validation", len(font_errors) > 0)
all_ok &= check("ASU does not auto-exclude unconfigured Heading 1 from body-spacing validation", len(spacing_errors) > 0)

raise SystemExit(0 if all_ok else 1)
