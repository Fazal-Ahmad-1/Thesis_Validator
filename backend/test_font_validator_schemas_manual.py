"""
Manual test for font_validator.py, covering both supported font rule
schemas (ASU's "allowed" list and Amity/Mumbai's "family" + "body_size").

Run from the backend/ directory with the venv active:
    python test_font_validator_schemas_manual.py

This is a manual test only (no pytest, no extra dependencies).
"""

from docx import Document

from validator.font_validator import validate_font_settings
from validator.rule_loader import load_rules


def make_doc(family: str, size_pt: float) -> Document:
    """Build a tiny in-memory document with one run in the given font."""
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Sample body text for font validation testing.")
    run.font.name = family
    from docx.shared import Pt
    run.font.size = Pt(size_pt)
    return document


def run_case(label, profile_name, family, size_pt, expect_pass):
    rules = load_rules(profile_name)
    document = make_doc(family, size_pt)
    errors = validate_font_settings(document, rules)

    passed = len(errors) == 0
    status = "PASS" if passed == expect_pass else "UNEXPECTED RESULT"

    print(f"[{status}] {label}")
    print(f"   profile={profile_name} font={family} {size_pt}pt -> "
          f"{'no violations' if passed else f'{len(errors)} violation(s)'}")
    if errors:
        for e in errors:
            print(f"     - {e['message']}")
    print()


def run_guard_check():
    """Confirm an unsupported font rule schema raises a clear ValueError."""
    document = make_doc("Times New Roman", 12)
    bad_rules = {"rules": {"font": {"something_else": True}}}
    print("Guard check: unsupported font rule schema...")
    try:
        validate_font_settings(document, bad_rules)
        print("UNEXPECTED: No error was raised.\n")
    except ValueError as e:
        print(f"Guard check passed. Raised ValueError as expected: {e}\n")


if __name__ == "__main__":
    # ASU: "allowed" list schema
    run_case("ASU valid font (Times New Roman 12pt, in allowed list)",
              "asu_graduate", "Times New Roman", 12, expect_pass=True)
    run_case("ASU invalid font (Calibri 11pt, not in allowed list)",
              "asu_graduate", "Calibri", 11, expect_pass=False)

    # Amity: "family" + "body_size" schema
    run_case("Amity valid font (Times New Roman 12pt)",
              "amity_selected", "Times New Roman", 12, expect_pass=True)
    run_case("Amity invalid font (Calibri 11pt)",
              "amity_selected", "Calibri", 11, expect_pass=False)

    # Mumbai: "family" + "body_size" schema
    run_case("Mumbai valid font (Times New Roman 12pt)",
              "mumbai_engineering", "Times New Roman", 12, expect_pass=True)
    run_case("Mumbai invalid font (Arial 10pt)",
              "mumbai_engineering", "Arial", 10, expect_pass=False)

    run_guard_check()