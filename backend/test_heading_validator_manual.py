"""
Manual test for heading_validator.py.

Run from the backend/ directory with the venv active:
    python test_heading_validator_manual.py

This is a manual test only (no pytest, no extra dependencies).
"""

from docx import Document
from docx.shared import Pt

from validator.heading_validator import validate_heading_settings
from validator.rule_loader import load_rules


def make_doc_with_heading(style_name, text, size_pt=None, bold=None, split_into_two_runs=False):
    """
    Build a tiny in-memory document with one heading paragraph.

    split_into_two_runs simulates a heading whose text got split across
    multiple runs (e.g. by a spell-checker), using the SAME formatting
    on both runs, to confirm grouping treats it as a single heading
    rather than multiple violations.
    """
    document = Document()
    paragraph = document.add_paragraph(style=style_name)

    if split_into_two_runs and len(text) > 1:
        midpoint = len(text) // 2
        pieces = [text[:midpoint], text[midpoint:]]
    else:
        pieces = [text]

    for piece in pieces:
        run = paragraph.add_run(piece)
        if size_pt is not None:
            run.font.size = Pt(size_pt)
        if bold is not None:
            run.font.bold = bold

    return document


def make_doc_with_mixed_runs(style_name, first_half, second_half, size_pt, bold_first, bold_second):
    """
    Build a heading paragraph where two runs have DIFFERENT bold values,
    to confirm this is reported as one "Mixed" violation, not two
    separate violations.
    """
    document = Document()
    paragraph = document.add_paragraph(style=style_name)

    run1 = paragraph.add_run(first_half)
    run1.font.size = Pt(size_pt)
    run1.font.bold = bold_first

    run2 = paragraph.add_run(second_half)
    run2.font.size = Pt(size_pt)
    run2.font.bold = bold_second

    return document


def run_case(label, document, rules, expect_error_count=None, expect_no_errors=None,
             expect_rule_contains=None, expect_actual=None):
    """
    expect_rule_contains: substring that must appear in every returned
        error's "rule" field (e.g. "Size", "Bold", "Uppercase").
    expect_actual: if given, asserts the single returned error's
        "actual" field equals this value (used for the Mixed case).
    """
    errors = validate_heading_settings(document, rules)

    if expect_no_errors:
        passed = len(errors) == 0
    else:
        passed = len(errors) == expect_error_count

    if passed and expect_rule_contains:
        passed = all(expect_rule_contains in e["rule"] for e in errors)

    if passed and expect_actual is not None:
        passed = len(errors) == 1 and errors[0]["actual"] == expect_actual

    status = "PASS" if passed else "UNEXPECTED RESULT"
    print(f"[{status}] {label} -> {len(errors)} error(s)")
    for e in errors:
        print(f"     - [{e['rule']}] actual={e['actual']!r} {e['message']} (occurrences={e['occurrences']})")
    print()


if __name__ == "__main__":
    amity_rules = load_rules("amity_selected")
    asu_rules = load_rules("asu_graduate")

    # --- Heading 1 (main_heading: 14pt, bold, uppercase) ---
    run_case(
        "Correctly formatted Heading 1",
        make_doc_with_heading("Heading 1", "INTRODUCTION", size_pt=14, bold=True),
        amity_rules,
        expect_error_count=0,
    )

    run_case(
        "Heading 1 with wrong size (12pt instead of 14pt)",
        make_doc_with_heading("Heading 1", "INTRODUCTION", size_pt=12, bold=True),
        amity_rules,
        expect_error_count=1,
        expect_rule_contains="Size",
    )

    run_case(
        "Heading 1 not bold",
        make_doc_with_heading("Heading 1", "INTRODUCTION", size_pt=14, bold=False),
        amity_rules,
        expect_error_count=1,
        expect_rule_contains="Bold",
    )

    run_case(
        "Heading 1 not uppercase",
        make_doc_with_heading("Heading 1", "Introduction", size_pt=14, bold=True),
        amity_rules,
        expect_error_count=1,
        expect_rule_contains="Uppercase",
    )

    run_case(
        "Heading 1 split into two runs, SAME wrong bold -> still ONE error (no duplicates)",
        make_doc_with_heading("Heading 1", "INTRODUCTION", size_pt=14, bold=False, split_into_two_runs=True),
        amity_rules,
        expect_error_count=1,
        expect_rule_contains="Bold",
    )

    run_case(
        "Heading 1 with two runs of DIFFERENT bold values -> ONE 'Mixed' violation",
        make_doc_with_mixed_runs("Heading 1", "INTRO", "DUCTION", size_pt=14, bold_first=True, bold_second=False),
        amity_rules,
        expect_error_count=1,
        expect_rule_contains="Bold",
        expect_actual="Mixed",
    )

    # --- Heading 2 (category_heading: 12pt, bold, uppercase) ---
    run_case(
        "Correctly formatted Heading 2",
        make_doc_with_heading("Heading 2", "LITERATURE REVIEW", size_pt=12, bold=True),
        amity_rules,
        expect_error_count=0,
    )

    run_case(
        "Heading 2 with wrong size (14pt instead of 12pt)",
        make_doc_with_heading("Heading 2", "LITERATURE REVIEW", size_pt=14, bold=True),
        amity_rules,
        expect_error_count=1,
        expect_rule_contains="Size",
    )

    run_case(
        "Heading 2 not bold",
        make_doc_with_heading("Heading 2", "LITERATURE REVIEW", size_pt=12, bold=False),
        amity_rules,
        expect_error_count=1,
        expect_rule_contains="Bold",
    )

    run_case(
        "Heading 2 not uppercase",
        make_doc_with_heading("Heading 2", "Literature Review", size_pt=12, bold=True),
        amity_rules,
        expect_error_count=1,
        expect_rule_contains="Uppercase",
    )

    # --- Heading 3 (sub_heading: 12pt, bold — no uppercase rule configured) ---
    run_case(
        "Correctly formatted Heading 3",
        make_doc_with_heading("Heading 3", "Background Studies", size_pt=12, bold=True),
        amity_rules,
        expect_error_count=0,
    )

    run_case(
        "Heading 3 with wrong size (10pt instead of 12pt)",
        make_doc_with_heading("Heading 3", "Background Studies", size_pt=10, bold=True),
        amity_rules,
        expect_error_count=1,
        expect_rule_contains="Size",
    )

    run_case(
        "Heading 3 not bold",
        make_doc_with_heading("Heading 3", "Background Studies", size_pt=12, bold=False),
        amity_rules,
        expect_error_count=1,
        expect_rule_contains="Bold",
    )

    run_case(
        "Heading 3 lowercase is NOT flagged (no uppercase rule configured for this level)",
        make_doc_with_heading("Heading 3", "background studies", size_pt=12, bold=True),
        amity_rules,
        expect_error_count=0,
    )

    # --- No headings in the document at all ---
    empty_body_doc = Document()
    empty_body_doc.add_paragraph("Just a normal paragraph, no headings at all.")
    run_case(
        "Document with no headings at all",
        empty_body_doc,
        amity_rules,
        expect_no_errors=True,
    )

    # --- Profile with no "headings" rules at all (ASU) ---
    run_case(
        "Profile without heading rules (ASU) always returns zero heading errors",
        make_doc_with_heading("Heading 1", "Anything Goes Here", size_pt=8, bold=False),
        asu_rules,
        expect_no_errors=True,
    )