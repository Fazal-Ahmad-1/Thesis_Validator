"""
Step 4A-2 reliability test: intentionally bad documents should be
detected as such by validate_document().

For each of the three rule profiles, this builds an in-memory DOCX with
deliberate violations in page size, margins, font, line spacing, and
(where required) a missing PAGE field. It then asserts:
    - is_valid is False
    - blocking_error_count > 0
    - score < 100
    - Page, Font, Paragraph, and Pagination all appear among the
      reported categories (Pagination only when the profile requires
      a PAGE field, which all three currently do)

Run from backend/ with the venv active:
    python test_reliability_bad_manual.py

Manual test only — no pytest, no new dependencies. Does not modify any
validator or rule file. Does not test headings or multiple sections
(out of scope for this step).
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, Mm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from validator.validator import validate_document


# --- Shared bad-formatting choices ---
# Deliberately wrong for every profile: Legal-ish oversized page, tiny
# uniform margins, an unapproved font/size, and single (1.0x) spacing
# when every current profile requires more than that.
BAD_PAGE_WIDTH = Inches(8.5)
BAD_PAGE_HEIGHT = Inches(14.0)   # Legal length — wrong for both Letter and A4
BAD_MARGIN = Inches(0.5)         # too small for every current profile
BAD_FONT_FAMILY = "Calibri"      # not in any current profile's allowed set
BAD_FONT_SIZE = 11
BAD_LINE_SPACING = 1.0           # every current profile requires 1.5 or 2.0


def add_bad_body_paragraph(document, text):
    """Add a body paragraph with explicit, deliberately wrong font/spacing."""
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = BAD_FONT_FAMILY
    run.font.size = Pt(BAD_FONT_SIZE)
    paragraph.paragraph_format.line_spacing = BAD_LINE_SPACING
    return paragraph


def apply_bad_page_settings(section):
    """Set deliberately wrong page size and margins on a section."""
    section.page_width = BAD_PAGE_WIDTH
    section.page_height = BAD_PAGE_HEIGHT
    section.top_margin = BAD_MARGIN
    section.bottom_margin = BAD_MARGIN
    section.left_margin = BAD_MARGIN
    section.right_margin = BAD_MARGIN


def build_bad_document():
    """
    Build one deliberately non-compliant document, reused across all
    three profiles: wrong page size/margins, wrong font, wrong line
    spacing, and NO PAGE field anywhere (footer left with plain text
    only, so the pagination check genuinely finds nothing).
    """
    document = Document()
    section = document.sections[0]

    apply_bad_page_settings(section)

    add_bad_body_paragraph(document, "Chapter 1: Introduction to the Study.")
    add_bad_body_paragraph(document, "This paragraph discusses the background of the research.")

    # Deliberately do NOT add a PAGE field. Put plain, non-field text in
    # the footer instead, so we're testing "field genuinely absent",
    # not "footer untouched" (which could be read as ambiguous).
    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.text = "Draft copy"

    return document


BUILDERS = {
    "asu_graduate": build_bad_document,
    "amity_selected": build_bad_document,
    "mumbai_engineering": build_bad_document,
}

EXPECTED_CATEGORIES = {"Page", "Font", "Paragraph", "Pagination"}


def run_case(profile_name):
    builder = BUILDERS[profile_name]
    document = builder()

    tmp_path = f"_tmp_bad_{profile_name}.docx"
    document.save(tmp_path)

    result = validate_document(tmp_path, profile_name)

    blocking_errors = result["blocking_errors"]
    detected_categories = {e.get("category") for e in blocking_errors if e.get("category")}

    checks = {
        "is_valid is False": result["is_valid"] is False,
        "blocking_error_count > 0": result["blocking_error_count"] > 0,
        "score < 100": result["score"] < 100,
        "all expected categories detected": EXPECTED_CATEGORIES.issubset(detected_categories),
    }

    passed = all(checks.values())
    status = "PASS" if passed else "FAIL"

    print(f"[{status}] {profile_name}")
    print(f"    is_valid={result['is_valid']}")
    print(f"    blocking_error_count={result['blocking_error_count']}")
    print(f"    score={result['score']}")
    print(f"    detected_categories={sorted(detected_categories)}")
    print(f"    expected_categories={sorted(EXPECTED_CATEGORIES)}")

    for check_name, ok in checks.items():
        if not ok:
            print(f"    FAILED CHECK: {check_name}")

    print("    Blocking errors:")
    for e in blocking_errors:
        print(f"      - [{e.get('category')}] {e.get('rule')}: {e.get('message')}")

    print()
    return passed


if __name__ == "__main__":
    print("Step 4A-2: intentionally bad document reliability check\n")

    results = {profile: run_case(profile) for profile in BUILDERS}

    print("=" * 60)
    for profile, passed in results.items():
        print(f"{profile}: {'PASS' if passed else 'FAIL'}")