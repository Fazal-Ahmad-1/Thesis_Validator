"""
Manual test script for paragraph_validator.py.

Run from the backend/ directory with the venv active:
    python test_paragraph_validator_manual.py

This is a manual test only (no pytest, no extra dependencies).
It does not modify any existing project files.
"""

from docx import Document

from validator.paragraph_validator import validate_paragraph_settings
from validator.rule_loader import load_rules

PROFILE_NAME = "asu_graduate"  # expects line_spacing == 2.0


def build_test_document() -> Document:
    """Create an in-memory DOCX document with paragraphs covering key cases."""
    document = Document()

    # Correct spacing (matches ASU's expected 2.0)
    p1 = document.add_paragraph("This paragraph has correct double spacing.")
    p1.paragraph_format.line_spacing = 2.0

    # Incorrect spacing: 1.5
    p2 = document.add_paragraph("This paragraph has incorrect 1.5 spacing.")
    p2.paragraph_format.line_spacing = 1.5

    # Incorrect spacing: 1.0 (first occurrence)
    p3 = document.add_paragraph("This paragraph has incorrect single spacing, first.")
    p3.paragraph_format.line_spacing = 1.0

    # Incorrect spacing: 1.0 (second occurrence, to test grouping/counting)
    p4 = document.add_paragraph("This paragraph has incorrect single spacing, second.")
    p4.paragraph_format.line_spacing = 1.0

    # Empty paragraph — should be ignored entirely
    document.add_paragraph("")

    return document


def run_paragraph_validation_test():
    """Build the test document, load rules, run validation, and print results."""
    document = build_test_document()
    rules = load_rules(PROFILE_NAME)

    errors = validate_paragraph_settings(document, rules)

    if not errors:
        print("No line spacing violations found.")
        return

    print(f"Found {len(errors)} unique line spacing violation(s):\n")

    for index, error in enumerate(errors, start=1):
        print(f"Violation {index}:")
        print(f"  rule:         {error['rule']}")
        print(f"  expected:     {error['expected']}")
        print(f"  actual:       {error['actual']}")
        print(f"  location:     {error['location']}")
        print(f"  text_snippet: {error['text_snippet']}")
        print(f"  occurrences:  {error['occurrences']}")
        print()

    print(f"Total number of unique violations: {len(errors)}")


if __name__ == "__main__":
    run_paragraph_validation_test()