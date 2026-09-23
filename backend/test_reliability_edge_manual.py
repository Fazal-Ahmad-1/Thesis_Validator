"""
Step 4B-1 reliability tests: mixed formatting and Word style inheritance.

Covers:
    1. Mixed font (some correct, some wrong) -> Font error expected.
    2. Mixed line spacing (some correct, some wrong) -> Paragraph error expected.
    3. Font inherited through the paragraph style's base_style chain,
       matching the profile -> no Font error expected.
    4. Line spacing inherited through the paragraph style's base_style
       chain, matching the profile -> no Paragraph error expected.
    5. Amity heading size/bold inherited from the built-in Heading 1/2/3
       Word styles (not set on individual runs) -> no Heading errors.

Run from backend/ with the venv active:
    python test_reliability_edge_manual.py

Manual test only — no pytest, no new dependencies. Does not modify any
validator or rule file.
"""

from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from validator.validator import validate_document

TMP_COUNTER = 0


# --- Shared OOXML helper (same proven pattern as prior reliability tests) ---

def add_page_field(paragraph):
    p = paragraph._p
    fld_simple = OxmlElement("w:fldSimple")
    fld_simple.set(qn("w:instr"), "PAGE")
    run_el = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = "1"
    run_el.append(text_el)
    fld_simple.append(run_el)
    p.append(fld_simple)


def make_asu_base_document():
    """
    A document that satisfies ASU's page/margin/pagination rules exactly,
    so each test below can isolate the ONE property it's actually
    exercising (font, or spacing) without unrelated Page/Pagination
    errors muddying the assertions.
    """
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer_paragraph)

    return document


def save_and_validate(document, profile_name):
    global TMP_COUNTER
    TMP_COUNTER += 1
    tmp_path = f"_tmp_edge_{TMP_COUNTER}.docx"
    document.save(tmp_path)
    return validate_document(tmp_path, profile_name)


def report(test_name, passed, result):
    categories = sorted({e.get("category") for e in result["blocking_errors"] if e.get("category")})
    print(f"[{'PASS' if passed else 'FAIL'}] {test_name}")
    print(f"    relevant_categories={categories}")
    print(f"    blocking_error_count={result['blocking_error_count']}")
    print(f"    score={result['score']}")
    if not passed:
        for e in result["blocking_errors"]:
            print(f"      - [{e.get('category')}] {e.get('rule')}: {e.get('message')}")
    print()
    return passed


# --- Test 1: mixed font ---

def test_mixed_font():
    document = make_asu_base_document()

    # Correct: Times New Roman 12pt (in ASU's allowed list)
    p1 = document.add_paragraph()
    r1 = p1.add_run("This paragraph uses the correct font.")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    p1.paragraph_format.line_spacing = 2.0

    # Wrong: Calibri 11pt (not in ASU's allowed list)
    p2 = document.add_paragraph()
    r2 = p2.add_run("This paragraph uses an incorrect font.")
    r2.font.name = "Calibri"
    r2.font.size = Pt(11)
    p2.paragraph_format.line_spacing = 2.0

    result = save_and_validate(document, "asu_graduate")
    categories = {e.get("category") for e in result["blocking_errors"]}

    passed = (
        "Font" in categories
        and result["is_valid"] is False
        and result["score"] < 100
    )
    return report("1. Mixed font (correct + incorrect)", passed, result)


# --- Test 2: mixed line spacing ---

def test_mixed_line_spacing():
    document = make_asu_base_document()

    # Correct: 2.0 spacing (ASU requirement), correct font throughout
    # so only the spacing violation is expected to fire.
    p1 = document.add_paragraph()
    r1 = p1.add_run("This paragraph uses the correct line spacing.")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    p1.paragraph_format.line_spacing = 2.0

    # Wrong: 1.0 spacing
    p2 = document.add_paragraph()
    r2 = p2.add_run("This paragraph uses incorrect line spacing.")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    p2.paragraph_format.line_spacing = 1.0

    result = save_and_validate(document, "asu_graduate")
    categories = {e.get("category") for e in result["blocking_errors"]}

    passed = (
        "Paragraph" in categories
        and result["is_valid"] is False
        and result["score"] < 100
    )
    return report("2. Mixed line spacing (correct + incorrect)", passed, result)


# --- Test 3: font inherited through the style hierarchy ---

def test_style_inherited_font():
    document = make_asu_base_document()

    # Base style declares the actual font; the style used by the
    # paragraph inherits from it via base_style and declares nothing
    # of its own. The run itself sets no font at all.
    base_style = document.styles.add_style("EdgeFontBase", WD_STYLE_TYPE.PARAGRAPH)
    base_style.font.name = "Times New Roman"
    base_style.font.size = Pt(12)

    derived_style = document.styles.add_style("EdgeFontDerived", WD_STYLE_TYPE.PARAGRAPH)
    derived_style.base_style = base_style

    p1 = document.add_paragraph(style="EdgeFontDerived")
    r1 = p1.add_run("This text's font comes entirely from the style chain.")
    # No run.font.name / run.font.size set intentionally.
    p1.paragraph_format.line_spacing = 2.0  # set directly so this test isolates font only

    result = save_and_validate(document, "asu_graduate")
    categories = {e.get("category") for e in result["blocking_errors"]}

    passed = "Font" not in categories
    return report("3. Font inherited via style base_style chain (matches profile)", passed, result)


# --- Test 4: line spacing inherited through the style hierarchy ---

def test_style_inherited_line_spacing():
    document = make_asu_base_document()

    base_style = document.styles.add_style("EdgeSpacingBase", WD_STYLE_TYPE.PARAGRAPH)
    base_style.paragraph_format.line_spacing = 2.0

    derived_style = document.styles.add_style("EdgeSpacingDerived", WD_STYLE_TYPE.PARAGRAPH)
    derived_style.base_style = base_style

    p1 = document.add_paragraph(style="EdgeSpacingDerived")
    r1 = p1.add_run("This paragraph's spacing comes entirely from the style chain.")
    r1.font.name = "Times New Roman"  # set directly so this test isolates spacing only
    r1.font.size = Pt(12)
    # No paragraph_format.line_spacing set directly on p1 intentionally.

    result = save_and_validate(document, "asu_graduate")
    categories = {e.get("category") for e in result["blocking_errors"]}

    passed = "Paragraph" not in categories
    return report("4. Line spacing inherited via style base_style chain (matches profile)", passed, result)


# --- Test 5: Amity heading size/bold inherited from Heading 1/2/3 styles ---

def test_amity_heading_style_inheritance():
    document = Document()
    section = document.sections[0]

    # Amity page settings (A4, 2.5cm margins) so Page checks don't fire.
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer_paragraph)

    # Set size/bold on the built-in Heading styles themselves, not on
    # any run, so heading_validator must resolve via paragraph.style.
    heading1_style = document.styles["Heading 1"]
    heading1_style.font.size = Pt(14)
    heading1_style.font.bold = True

    heading2_style = document.styles["Heading 2"]
    heading2_style.font.size = Pt(12)
    heading2_style.font.bold = True

    heading3_style = document.styles["Heading 3"]
    heading3_style.font.size = Pt(12)
    heading3_style.font.bold = True

    p1 = document.add_paragraph(style="Heading 1")
    p1.add_run("MAIN HEADING")  # uppercase text set directly; no run.font set

    p2 = document.add_paragraph(style="Heading 2")
    p2.add_run("CATEGORY HEADING")

    p3 = document.add_paragraph(style="Heading 3")
    p3.add_run("Sub Heading")

    result = save_and_validate(document, "amity_selected")
    heading_errors = [e for e in result["blocking_errors"] if e.get("category") == "Heading"]

    passed = len(heading_errors) == 0
    return report("5. Amity H1/H2/H3 size+bold inherited from Word styles", passed, result)


if __name__ == "__main__":
    print("Step 4B-1: mixed formatting and style inheritance reliability checks\n")

    results = {
        "mixed_font": test_mixed_font(),
        "mixed_line_spacing": test_mixed_line_spacing(),
        "style_inherited_font": test_style_inherited_font(),
        "style_inherited_line_spacing": test_style_inherited_line_spacing(),
        "amity_heading_style_inheritance": test_amity_heading_style_inheritance(),
    }

    print("=" * 60)
    for name, passed in results.items():
        print(f"{name}: {'PASS' if passed else 'FAIL'}")